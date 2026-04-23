import os
import zipfile
import shutil
from flask import Flask, render_template, request, redirect, send_file
from Sastrawi.Stemmer.StemmerFactory import StemmerFactory
import PyPDF2
import docx
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from fpdf import FPDF

app = Flask(__name__)

# --- KONFIGURASI ---
basedir = os.path.abspath(os.path.dirname(__file__))
app.config['UPLOAD_FOLDER'] = os.path.join(basedir, 'uploads')
app.config['MAX_CONTENT_LENGTH'] = 32 * 1024 * 1024

# Inisialisasi Sastrawi
factory = StemmerFactory()
stemmer = factory.create_stemmer()

# Simpan hasil terakhir untuk export PDF
last_results = []

def ekstrak_teks(file_path, filename):
    teks = ""
    ext = filename.split('.')[-1].lower()
    try:
        if ext == 'pdf':
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    teks += page.extract_text() or ""
        elif ext == 'docx':
            doc = docx.Document(file_path)
            teks = " ".join([p.text for p in doc.paragraphs])
        elif ext == 'txt':
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                teks = f.read()
    except:
        pass
    return teks

@app.route('/')
def index():
    return render_template('index.html', page='home')

@app.route('/upload', methods=['POST'])
def upload_file():
    global last_results
    files = request.files.getlist('file_tugas')
    threshold = float(request.form.get('threshold', 25))
    
    if not files or files[0].filename == '':
        return redirect('/')

    if not os.path.exists(app.config['UPLOAD_FOLDER']):
        os.makedirs(app.config['UPLOAD_FOLDER'])

    data_sesi_ini = []
    for file in files:
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
        file.save(file_path)

        if file.filename.endswith('.zip'):
            folder_ekstrak = os.path.join(app.config['UPLOAD_FOLDER'], 'temp_' + file.filename)
            with zipfile.ZipFile(file_path, 'r') as zip_ref:
                zip_ref.extractall(folder_ekstrak)
            for root, _, f_names in os.walk(folder_ekstrak):
                for f_name in f_names:
                    if f_name.split('.')[-1].lower() in ['pdf', 'docx', 'txt']:
                        p_path = os.path.join(root, f_name)
                        konten = ekstrak_teks(p_path, f_name)
                        if konten.strip():
                            data_sesi_ini.append({'judul': f_name, 'konten': stemmer.stem(konten)})
            shutil.rmtree(folder_ekstrak)
        else:
            konten = ekstrak_teks(file_path, file.filename)
            if konten.strip():
                data_sesi_ini.append({'judul': file.filename, 'konten': stemmer.stem(konten)})

    hasil_final = []
    if len(data_sesi_ini) > 1:
        korpus = [d['konten'] for d in data_sesi_ini]
        tfidf_matrix = TfidfVectorizer().fit_transform(korpus)
        sim_matrix = cosine_similarity(tfidf_matrix)

        for i in range(len(data_sesi_ini)):
            for j in range(i + 1, len(data_sesi_ini)):
                skor = round(sim_matrix[i][j] * 100, 2)
                if skor >= threshold:
                    hasil_final.append({
                        'doc1': data_sesi_ini[i]['judul'],
                        'doc2': data_sesi_ini[j]['judul'],
                        'skor': skor
                    })
    
    last_results = hasil_final
    return render_template('index.html', page='result', hasil=hasil_final, threshold=threshold)

@app.route('/export')
def export_pdf():
    pdf = FPDF()
    pdf.add_page()
    
    # Header Laporan
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(190, 10, txt="LAPORAN HASIL ANALISIS PLAGIARISME", ln=True, align='C')
    pdf.set_font("Arial", '', 10)
    pdf.cell(190, 10, txt="Metode: Cosine Similarity & TF-IDF", ln=True, align='C')
    pdf.ln(10)
    
    # Header Tabel
    pdf.set_font("Arial", 'B', 11)
    pdf.set_fill_color(230, 230, 230)
    
    w_doc = 75    # Lebar kolom dokumen
    w_score = 40  # Lebar kolom skor
    h_min = 10    # Tinggi baris minimal
    
    pdf.cell(w_doc, h_min, " Dokumen A", 1, 0, 'L', True)
    pdf.cell(w_doc, h_min, " Dokumen B", 1, 0, 'L', True)
    pdf.cell(w_score, h_min, "Skor (%)", 1, 1, 'C', True)
    
    pdf.set_font("Arial", '', 9)
    
    for res in last_results:
        # 1. Simpan posisi awal Y
        y_start = pdf.get_y()
        x_start = pdf.get_x()

        # 2. Hitung tinggi yang dibutuhkan (Simulasi penulisan)
        # Kita hitung jumlah baris teks jika di-wrap di lebar w_doc
        lines_a = len(pdf.multi_cell(w_doc, 5, res['doc1'], split_only=True))
        lines_b = len(pdf.multi_cell(w_doc, 5, res['doc2'], split_only=True))
        max_lines = max(lines_a, lines_b)
        
        # Tentukan tinggi baris final (misal 1 baris teks = 6 unit)
        calculated_height = max_lines * 6
        final_h = max(h_min, calculated_height)

        # 3. Gambar kolom A
        pdf.multi_cell(w_doc, final_h / lines_a if lines_a > 0 else final_h, res['doc1'], 1, 'L')
        
        # 4. Paksa posisi Y kembali ke awal baris untuk kolom B
        pdf.set_xy(x_start + w_doc, y_start)
        pdf.multi_cell(w_doc, final_h / lines_b if lines_b > 0 else final_h, res['doc2'], 1, 'L')
        
        # 5. Paksa posisi Y kembali ke awal baris untuk kolom Skor
        pdf.set_xy(x_start + (w_doc * 2), y_start)
        # Gunakan cell biasa agar alignment center secara vertikal sempurna
        pdf.cell(w_score, final_h, f"{res['skor']}%", 1, 1, 'C')
        
        # 6. Set Y ke posisi paling bawah dari baris ini sebelum lanjut ke baris berikutnya
        pdf.set_y(y_start + final_h)

    path = os.path.join(app.config['UPLOAD_FOLDER'], "laporan_plagiarisme_final.pdf")
    pdf.output(path)
    return send_file(path, as_attachment=True)

@app.route('/about')
def about():
    return render_template('index.html', page='about')

if __name__ == '__main__':
    app.run(debug=True)